import csv
import json
from decimal import Decimal
from functools import wraps
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.utils import timezone

from .forms import ExpenseFormSet, NewConvertFormSet, NewcomerFormSet, ServiceReportForm, ExtraIncomeForm, ExtraExpenseForm
from .forms import ReportIncomeLineFormSet
from .models import ReportIncomeLine
from .models import Expense, ExtraIncome, ExtraExpense, ServiceReport
from .permissions import reports_for_user, user_extension, user_is_admin
from .realtime import publish_report_created
from .services import money, prepare_income_lines
from apps.churches.currency_service import convert_currency, get_currency_for_extension, get_record_currency, CurrencyConversionError
from apps.churches.models import Currency

MOIS_FR = [
    "", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
    "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre",
]


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def conversion_required(view):
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        try:
            return view(request, *args, **kwargs)
        except CurrencyConversionError as exc:
            return render(request, "reports/conversion_unavailable.html", {"reason": str(exc)}, status=422)
    return wrapped


def _report_filename(report, extension):
    date_part = report.service_date.isoformat()
    slug = report.extension.slug
    return f"rapport-{slug}-{date_part}.{extension}"


def _money(value):
    return f"{value:.2f}"


def _report_rows(report):
    rows = [
        ("Extension", report.extension.name),
        ("Date", report.service_date),
        ("Devise", get_record_currency(report).code if get_record_currency(report) else "Non renseignée"),
        ("Type", report.get_service_type_display()),
        ("Prédicateur", report.preacher or ""),
        ("Modérateur", report.moderator or ""),
        ("Interprète", report.interpreter or ""),
        ("Textes bibliques", report.scripture_text or ""),
        ("Thème", report.theme or ""),
        ("Papa", report.papa_count),
        ("Maman", report.maman_count),
        ("Frères", report.brothers_count),
        ("Sœurs", report.sisters_count),
        ("Enfants", report.children_count),
        ("Présence totale", report.total_attendance),
        ("Offrandes ordinaires", _money(report.offering_regular)),
        ("Offrandes orateur", _money(report.offering_preacher)),
        ("Dîmes", _money(report.offering_tithe)),
        ("Actions de grâce", _money(report.offering_thanksgiving)),
        ("Total offrandes", _money(report.total_offerings)),
        ("Prélèvement dîmes", _money(report.tithe_deduction)),
        ("Prélèvement social", _money(report.social_deduction)),
        ("Dépenses du culte", _money(report.total_expenses)),
        ("Solde net du culte", _money(report.net_balance)),
    ]

    for line in report.income_lines.select_related("currency"):
        rows.append(("Entrée — " + line.get_category_display(),
                     f"{line.amount} {line.currency.code} × {line.exchange_rate} = {line.converted_amount} {get_record_currency(report).code}"))
    return rows


def _get_ventilation(report):
    """Retourne la ventilation ligne par ligne pour l'affichage du rapport.

    Utilise les taux figés (snapshots) enregistrés lors de la création du
    rapport, de sorte que l'historique reste stable même après une modification
    des paramètres de l'extension.
    """
    from .services import compute_ventilation
    return compute_ventilation(
        report.offering_regular,
        report.offering_preacher,
        report.offering_tithe,
        report.offering_thanksgiving,
        report.tithe_percentage_snapshot,
        report.social_percentage_snapshot,
    )


@login_required
def report_financial_preview(request):
    """Endpoint JSON utilisé par le formulaire pour le calcul temps réel.

    Attend un POST JSON avec les champs :
      - extension_id (int)
      - offering_regular, offering_preacher, offering_tithe, offering_thanksgiving (Decimal)
      - expenses : liste de montants numériques

    Retourne la ventilation complète et le reste final.
    """
    from apps.churches.models import ChurchExtension
    from .services import compute_ventilation, money

    if request.method != "POST":
        return JsonResponse({"error": "Méthode non autorisée"}, status=405)

    try:
        data = json.loads(request.body)
    except (json.JSONDecodeError, ValueError):
        return JsonResponse({"error": "JSON invalide"}, status=400)

    def to_decimal(val):
        try:
            return money(max(Decimal("0"), Decimal(str(val or 0))))
        except Exception:
            return Decimal("0")

    # Récupérer les taux de l'extension
    tithe_pct  = Decimal("10")
    social_pct = Decimal("10")
    extension_id = data.get("extension_id")
    if extension_id:
        try:
            ext = ChurchExtension.objects.get(pk=int(extension_id))
            tithe_pct  = ext.tithe_percentage
            social_pct = ext.social_percentage
        except (ChurchExtension.DoesNotExist, ValueError, TypeError):
            pass

    ordinaires    = to_decimal(data.get("offering_regular",       0))
    orateur       = to_decimal(data.get("offering_preacher",      0))
    dimes         = to_decimal(data.get("offering_tithe",         0))
    actions_grace = to_decimal(data.get("offering_thanksgiving",  0))
    try:
        target = Currency.objects.filter(pk=data.get("currency_id")).first() if data.get("currency_id") else None
        if not target:
            preview_ext = ChurchExtension.objects.filter(pk=extension_id).first() if extension_id else user_extension(request.user)
            target = get_currency_for_extension(preview_ext)
        raw_lines = data.get("income_lines", [])
        if not isinstance(raw_lines, list) or len(raw_lines) > 100:
            raise ValueError("Nombre de lignes invalide")
        lines = []
        for item in raw_lines:
            amount = Decimal(str(item["amount"]))
            if not amount.is_finite() or amount <= 0 or item["category"] not in ReportIncomeLine.Category.values:
                raise ValueError("Entrée invalide")
            currency = Currency.objects.filter(pk=item["currency"], is_active=True).first()
            if not currency:
                raise ValueError("Devise invalide")
            lines.append(ReportIncomeLine(category=item["category"], amount=amount, currency=currency))
        converted, _ = prepare_income_lines({
            "offering_regular": ordinaires, "offering_preacher": orateur,
            "offering_tithe": dimes, "offering_thanksgiving": actions_grace,
        }, lines, target)
        ordinaires, orateur, dimes, actions_grace = [converted[key] for key in ReportIncomeLine.Category.values]
    except (CurrencyConversionError, ValueError, TypeError, KeyError, ArithmeticError) as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    expenses_raw  = data.get("expenses", [])

    total_expenses = money(sum(
        to_decimal(e) for e in expenses_raw if e is not None
    ))

    v = compute_ventilation(ordinaires, orateur, dimes, actions_grace, tithe_pct, social_pct)

    total_recettes = money(ordinaires + orateur + dimes + actions_grace)
    reste_final    = money(v["offering_remainder"] - total_expenses)

    def fmt(d):
        return str(d)

    return JsonResponse({
        "currency_code": target.code if target else "",
        "amounts": {key: str(value) for key, value in converted.items()},
        "tithe_pct":   str(tithe_pct),
        "social_pct":  str(social_pct),
        "rows": {
            "ordinaires":    {"dime": fmt(v["ordinaires"]["dime"]),    "social": fmt(v["ordinaires"]["social"]),    "reste": fmt(v["ordinaires"]["reste"])},
            "orateur":       {"dime": fmt(v["orateur"]["dime"]),       "social": fmt(v["orateur"]["social"]),       "reste": fmt(v["orateur"]["reste"])},
            "dimes":         {"dime": fmt(v["dimes"]["dime"]),         "social": fmt(v["dimes"]["social"]),         "reste": fmt(v["dimes"]["reste"])},
            "actions_grace": {"dime": fmt(v["actions_grace"]["dime"]), "social": fmt(v["actions_grace"]["social"]), "reste": fmt(v["actions_grace"]["reste"])},
        },
        "totals": {
            "recettes":  fmt(total_recettes),
            "dime":      fmt(v["total_dime"]),
            "social":    fmt(v["total_social"]),
            "reste":     fmt(v["offering_remainder"]),
            "depenses":  fmt(total_expenses),
            "reste_final": fmt(reste_final),
        },
    })


# ─────────────────────────────────────────────────────────────────────────────
# Export CSV / DOCX / PDF
# ─────────────────────────────────────────────────────────────────────────────

def _export_report_csv(report):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_report_filename(report, "csv")}"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(["Champ", "Valeur"])
    writer.writerows(_report_rows(report))
    return response


def _export_report_docx(report):
    from docx import Document
    from docx.shared import Inches, RGBColor
    from django.contrib.staticfiles import finders

    document = Document()
    document.add_picture(finders.find("branding/emerge-symbol.png"), width=Inches(0.65))
    document.add_paragraph("Emerge In Christ - Rassemblement")
    document.styles["Heading 1"].font.color.rgb = RGBColor.from_string("C90800")
    document.add_heading(f"Rapport — {report.extension.name}", level=1)
    document.add_paragraph(f"{report.get_service_type_display()} du {report.service_date}")

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Champ"
    table.rows[0].cells[1].text = "Valeur"
    for label, value in _report_rows(report):
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)

    buffer = BytesIO()
    document.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = f'attachment; filename="{_report_filename(report, "docx")}"'
    return response


def _export_report_pdf(report):
    from xml.sax.saxutils import escape
    from django.contrib.staticfiles import finders
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph, Image

    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, title=f"Rapport {report.extension.name}")
    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.HexColor("#C90800")
    styles["Heading2"].textColor = colors.HexColor("#9A642E")
    rows = [["Champ", "Valeur"], *[
        [Paragraph(escape(str(label)), styles["BodyText"]),
         Paragraph(escape(str(value)), styles["BodyText"])]
        for label, value in _report_rows(report)
    ]]
    table = Table(rows, colWidths=[150, document.width - 150], repeatRows=1)
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C90800")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
        ])
    )
    story = [
        Image(finders.find("branding/emerge-symbol.png"), width=50, height=50 * 622 / 531),
        Spacer(1, 8),
        Paragraph("Emerge In Christ - Rassemblement", styles["Heading2"]),
        Paragraph(f"Rapport - {escape(report.extension.name)}", styles["Title"]),
        Paragraph(f"{report.get_service_type_display()} du {report.service_date}", styles["Normal"]),
        Spacer(1, 18),
        table,
    ]
    document.build(story)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_report_filename(report, "pdf")}"'
    return response


def _export_reports_csv(reports):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="rapports.csv"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow([
        "Date", "Extension", "Type", "Présence",
        "Offrandes", "Dîmes", "Social", "Dépenses", "Solde", "Devise",
    ])
    for report in reports:
        writer.writerow([
            report.service_date,
            report.extension.name,
            report.get_service_type_display(),
            report.total_attendance,
            _money(report.total_offerings),
            _money(report.tithe_deduction),
            _money(report.social_deduction),
            _money(report.total_expenses),
            _money(report.net_balance),
            get_record_currency(report).code if get_record_currency(report) else "Non renseignée",
        ])
    return response


# ─────────────────────────────────────────────────────────────────────────────
# Vues principales
# ─────────────────────────────────────────────────────────────────────────────

@login_required
def report_list(request):
    reports = reports_for_user(request.user)
    return render(request, "reports/report_list.html", {"reports": reports})


@login_required
def report_detail(request, pk):
    report = get_object_or_404(reports_for_user(request.user), pk=pk)
    ventilation = _get_ventilation(report)
    ext = report.extension
    ext_currency = get_record_currency(report)
    try:
        usd_balance = convert_currency(report.net_balance, ext_currency, "USD")
    except CurrencyConversionError:
        usd_balance = None
    return render(request, "reports/report_detail.html", {
        "report": report,
        "total_deductions": report.tithe_deduction + report.social_deduction,
        "ventilation": ventilation,
        "income_lines": report.income_lines.select_related("currency"),
        "expense_items": report.expense_items.all(),
        "newcomers": report.newcomers.all(),
        "new_converts": report.new_converts.all(),
        "usd_balance": usd_balance,
        "show_usd": bool(ext_currency and ext_currency.code != "USD" and usd_balance is not None),
        "report_currency": ext_currency,
    })


@login_required
def report_print(request, pk):
    report = get_object_or_404(reports_for_user(request.user), pk=pk)
    ventilation = _get_ventilation(report)
    ext = report.extension
    ext_currency = get_record_currency(report)
    try:
        usd_balance = convert_currency(report.net_balance, ext_currency, "USD")
    except CurrencyConversionError:
        usd_balance = None
    return render(request, "reports/report_print.html", {
        "report": report,
        "ventilation": ventilation,
        "income_lines": report.income_lines.select_related("currency"),
        "expense_items": report.expense_items.all(),
        "newcomers": report.newcomers.all(),
        "new_converts": report.new_converts.all(),
        "usd_balance": usd_balance,
        "show_usd": bool(ext_currency and ext_currency.code != "USD" and usd_balance is not None),
        "report_currency": ext_currency,
        "currency_symbol": ext_currency.symbol if ext_currency else "—",
        "currency_code": ext_currency.code if ext_currency else "Non renseignée",
        "exchange_rate_to_usd": ext_currency.usd_rate if ext_currency else 1,
    })


@login_required
def report_export(request, pk, file_format):
    report = get_object_or_404(reports_for_user(request.user), pk=pk)
    exporters = {
        "csv": _export_report_csv,
        "docx": _export_report_docx,
        "pdf": _export_report_pdf,
    }
    exporter = exporters.get(file_format)
    if exporter is None:
        return redirect("reports:detail", pk=report.pk)
    return exporter(report)


@login_required
@conversion_required
def reports_export(request, file_format):
    reports = reports_for_user(request.user)
    if file_format == "csv":
        return _export_reports_csv(reports)
    if file_format == "html":
        target = Currency.get_default() if user_is_admin(request.user) else get_currency_for_extension(user_extension(request.user))
        reports = list(reports.select_related("currency", "extension__currency"))
        fields = ["total_offerings", "tithe_deduction", "social_deduction", "net_balance"]
        totals = {field: Decimal("0.00") for field in fields}
        totals["total_attendance"] = 0
        for report in reports:
            totals["total_attendance"] += report.total_attendance
            for field in fields:
                amount = convert_currency(getattr(report, field), get_record_currency(report), target)
                setattr(report, field + "_disp", amount)
                totals[field] += amount
        html = render_to_string(
            "reports/report_collection_print.html", {"reports": reports, "totals": totals, "display_currency": target}, request=request
        )
        return HttpResponse(html)
    return redirect("reports:list")


@login_required
@transaction.atomic
def report_create(request):
    initial = {}
    if not user_is_admin(request.user):
        initial["extension"] = user_extension(request.user)

    form = ServiceReportForm(request.POST or None, initial=initial)
    income_formset = ReportIncomeLineFormSet(request.POST or None, prefix="income_lines")
    expense_formset = ExpenseFormSet(request.POST or None, prefix="expenses")
    newcomer_formset = NewcomerFormSet(request.POST or None, prefix="newcomers")
    new_convert_formset = NewConvertFormSet(request.POST or None, prefix="converts")

    if not user_is_admin(request.user):
        form.fields["extension"].disabled = True
        form.fields["extension"].required = False

    valid = all([form.is_valid(), expense_formset.is_valid(), newcomer_formset.is_valid(), new_convert_formset.is_valid(), income_formset.is_valid()])
    prepared_lines = []
    if valid:
        try:
            converted_totals, prepared_lines = prepare_income_lines(form.cleaned_data, income_formset.save(commit=False), form.cleaned_data["currency"])
        except CurrencyConversionError as exc:
            form.add_error(None, str(exc))
            valid = False
    if valid:
        report = form.save(commit=False)
        for field, amount in converted_totals.items():
            setattr(report, field, amount)
        if not user_is_admin(request.user):
            report.extension = user_extension(request.user)
        report.submitted_by = request.user
        report.save()
        for line in prepared_lines:
            line.report = report
            line.save()

        # Sauvegarder les lignes de dépenses
        expenses = expense_formset.save(commit=False)
        for i, expense in enumerate(expenses, start=1):
            expense.report = report
            expense.number = i
            expense.save()
        for obj in expense_formset.deleted_objects:
            obj.delete()

        # Sauvegarder les nouveaux venus
        newcomers = newcomer_formset.save(commit=False)
        for newcomer in newcomers:
            newcomer.report = report
            newcomer.save()
        for obj in newcomer_formset.deleted_objects:
            obj.delete()

        # Sauvegarder les ames gagnees
        converts = new_convert_formset.save(commit=False)
        for convert in converts:
            convert.report = report
            convert.save()
        for obj in new_convert_formset.deleted_objects:
            obj.delete()

        # Recalculer les totaux après les lignes
        report.refresh_totals()

        # Notifications & email
        try:
            from apps.notifications.services import notify_all_admins
            from apps.notifications.models import Notification
            from apps.notifications.email_service import send_report_submitted_email

            notify_all_admins(
                title=f"Nouveau rapport — {report.extension.name}",
                body=(
                    f"Un nouveau rapport de culte a été soumis par "
                    f"{request.user.get_full_name() or request.user.username} "
                    f"pour {report.extension.name} le {report.service_date}."
                ),
                notification_type=Notification.Type.REPORT_CREATED,
                link=f"/rapports/{report.pk}/",
            )
            send_report_submitted_email(report)
        except Exception:
            pass  # Ne pas bloquer la création du rapport si l'email échoue

        publish_report_created(report)
        return redirect("reports:detail", pk=report.pk)

    return render(request, "reports/report_form.html", {
        "form": form,
        "income_formset": income_formset,
        "expense_formset": expense_formset,
        "newcomer_formset": newcomer_formset,
        "new_convert_formset": new_convert_formset,
    })


# ─────────────────────────────────────────────────────────────────────────────
# Recettes et Dépenses supplémentaires
# ─────────────────────────────────────────────────────────────────────────────

@login_required
def extra_income_list(request):
    incomes = ExtraIncome.objects.all().order_by("-income_date")
    if not user_is_admin(request.user):
        incomes = incomes.filter(extension=user_extension(request.user))
    return render(request, "reports/extra_income_list.html", {"incomes": incomes})

@login_required
def extra_income_create(request):
    initial = {}
    if not user_is_admin(request.user):
        initial["extension"] = user_extension(request.user)

    form = ExtraIncomeForm(request.POST or None, initial=initial)
    if not user_is_admin(request.user):
        form.fields["extension"].disabled = True
        form.fields["extension"].required = False

    if form.is_valid():
        income = form.save(commit=False)
        if not user_is_admin(request.user):
            income.extension = user_extension(request.user)
        income.save()
        return redirect("reports:extra_income_list")

    return render(request, "reports/extra_income_form.html", {"form": form})

@login_required
def extra_expense_list(request):
    expenses = ExtraExpense.objects.all().order_by("-expense_date")
    if not user_is_admin(request.user):
        expenses = expenses.filter(extension=user_extension(request.user))
    return render(request, "reports/extra_expense_list.html", {"expenses": expenses})

@login_required
def extra_expense_create(request):
    initial = {}
    if not user_is_admin(request.user):
        initial["extension"] = user_extension(request.user)

    form = ExtraExpenseForm(request.POST or None, initial=initial)
    if not user_is_admin(request.user):
        form.fields["extension"].disabled = True
        form.fields["extension"].required = False

    if form.is_valid():
        expense = form.save(commit=False)
        if not user_is_admin(request.user):
            expense.extension = user_extension(request.user)
        expense.save()
        return redirect("reports:extra_expense_list")

    return render(request, "reports/extra_expense_form.html", {"form": form})


# ─────────────────────────────────────────────────────────────────────────────
# Rapports périodiques
# ─────────────────────────────────────────────────────────────────────────────

@login_required
@conversion_required
def report_monthly(request, year, month):
    """Vue du rapport mensuel (affichage dans le navigateur)."""
    extension_id = request.GET.get("extension")
    reports = reports_for_user(request.user).filter(
        service_date__year=year, service_date__month=month
    ).select_related("extension")
    
    totals, selected_ext, is_consolidated, reports_filtered, extras_income, extras_expense = _aggregate_totals(
        reports, year, month=month, user=request.user, extension_id=extension_id
    )
    
    return render(request, "reports/report_monthly.html", {
        "reports": reports_filtered,
        "totals": totals,
        "year": year,
        "month": month,
        "month_name": MOIS_FR[month],
    })


@login_required
@conversion_required
def report_monthly_print(request, year, month):
    """Rapport mensuel version imprimable."""
    extension_id = request.GET.get("extension")
    reports = reports_for_user(request.user).filter(
        service_date__year=year, service_date__month=month
    ).select_related("extension")
    
    totals, selected_ext, is_consolidated, reports_filtered, extras_income, extras_expense = _aggregate_totals(
        reports, year, month=month, user=request.user, extension_id=extension_id
    )
    
    ext = selected_ext
    ext_currency = get_currency_for_extension(ext) if ext else Currency.get_default()
    sym = ext_currency.symbol if ext_currency else "$"
    
    usd_totals = _totals_in_usd(totals, ext) if (ext and not is_consolidated) else None
    
    from apps.churches.models import ChurchExtension
    extensions = ChurchExtension.objects.filter(is_active=True) if user_is_admin(request.user) else None
    
    return render(request, "reports/report_monthly_print.html", {
        "reports": reports_filtered,
        "totals": totals,
        "usd_totals": usd_totals,
        "year": year,
        "month": month,
        "month_name": MOIS_FR[month],
        "extension": ext,
        "currency_symbol": sym,
        "currency_code": ext_currency.code if ext_currency else "Non renseignée",
        "exchange_rate_to_usd": ext_currency.usd_rate if ext_currency else 1,
        "show_usd": bool(ext_currency and ext_currency.code != "USD" and not is_consolidated),
        "is_consolidated": is_consolidated,
        "extensions": extensions,
        "is_admin": user_is_admin(request.user),
        "extras_income": extras_income,
        "extras_expense": extras_expense,
    })


@login_required
@conversion_required
def report_quarterly_print(request, year, quarter):
    """Rapport trimestriel version imprimable. quarter = 1,2,3,4"""
    quarter = int(quarter)
    months = {1: [1, 2, 3], 2: [4, 5, 6], 3: [7, 8, 9], 4: [10, 11, 12]}
    quarter_months = months.get(quarter, [1, 2, 3])
    extension_id = request.GET.get("extension")

    # D'abord calculer les totaux trimestriels globaux
    all_reports = reports_for_user(request.user).filter(
        service_date__year=year,
        service_date__month__in=quarter_months,
    )
    
    quarter_totals, selected_ext, is_consolidated, all_reports_filtered, extras_income, extras_expense = _aggregate_totals(
        all_reports, year, months=quarter_months, user=request.user, extension_id=extension_id
    )
    
    # Ensuite calculer les données mensuelles en réutilisant les rapports filtrés
    month_data = []
    for m in quarter_months:
        # Filtrer les rapports déjà récupérés pour ce mois spécifique
        reports_m = [r for r in all_reports_filtered if r.service_date.month == m]
        
        # Recalculer les totaux pour ce mois uniquement
        totals_m, _, _, _, _, _ = _aggregate_totals(
            reports_m, year, month=m, user=request.user, extension_id=extension_id
        )
        
        month_data.append({
            "month": m,
            "month_name": MOIS_FR[m],
            "reports": reports_m,
            "totals": totals_m,
        })
    
    ext = selected_ext
    ext_currency = get_currency_for_extension(ext) if ext else Currency.get_default()
    sym = ext_currency.symbol if ext_currency else "$"
    
    usd_totals = _totals_in_usd(quarter_totals, ext) if (ext and not is_consolidated) else None
    
    from apps.churches.models import ChurchExtension
    extensions = ChurchExtension.objects.filter(is_active=True) if user_is_admin(request.user) else None
    
    quarter_names = {1: "1er Trimestre", 2: "2ème Trimestre",
                     3: "3ème Trimestre", 4: "4ème Trimestre"}

    return render(request, "reports/report_quarterly_print.html", {
        "month_data": month_data,
        "quarter_totals": quarter_totals,
        "usd_totals": usd_totals,
        "year": year,
        "quarter": quarter,
        "quarter_name": quarter_names.get(quarter, f"T{quarter}"),
        "extension": ext,
        "currency_symbol": sym,
        "currency_code": ext_currency.code if ext_currency else "Non renseignée",
        "exchange_rate_to_usd": ext_currency.usd_rate if ext_currency else 1,
        "show_usd": bool(ext_currency and ext_currency.code != "USD" and not is_consolidated),
        "is_consolidated": is_consolidated,
        "extensions": extensions,
        "is_admin": user_is_admin(request.user),
        "extras_income": extras_income,
        "extras_expense": extras_expense,
    })


@login_required
@conversion_required
def report_annual_print(request, year):
    """Rapport annuel version imprimable."""
    extension_id = request.GET.get("extension")
    
    # D'abord calculer les totaux annuels globaux
    all_reports = reports_for_user(request.user).filter(service_date__year=year)
    
    annual_totals, selected_ext, is_consolidated, all_reports_filtered, extras_income, extras_expense = _aggregate_totals(
        all_reports, year, user=request.user, extension_id=extension_id
    )
    
    # Ensuite calculer les données trimestrielles en réutilisant les rapports filtrés
    quarter_data = []
    quarter_names = {1: "1er Trimestre", 2: "2ème Trimestre",
                     3: "3ème Trimestre", 4: "4ème Trimestre"}
    for q, months in {1: [1,2,3], 2:[4,5,6], 3:[7,8,9], 4:[10,11,12]}.items():
        # Filtrer les rapports déjà récupérés pour ce trimestre spécifique
        reports_q = [r for r in all_reports_filtered if r.service_date.month in months]
        
        # Recalculer les totaux pour ce trimestre uniquement
        totals_q, _, _, _, _, _ = _aggregate_totals(
            reports_q, year, months=months, user=request.user, extension_id=extension_id
        )
        
        quarter_data.append({
            "quarter": q,
            "quarter_name": quarter_names[q],
            "reports": reports_q,
            "totals": totals_q,
            "months": [{"month": m, "month_name": MOIS_FR[m]} for m in months],
        })
    
    ext = selected_ext
    ext_currency = get_currency_for_extension(ext) if ext else Currency.get_default()
    sym = ext_currency.symbol if ext_currency else "$"
    
    usd_totals = _totals_in_usd(annual_totals, ext) if (ext and not is_consolidated) else None
    
    from apps.churches.models import ChurchExtension
    extensions = ChurchExtension.objects.filter(is_active=True) if user_is_admin(request.user) else None

    return render(request, "reports/report_annual_print.html", {
        "quarter_data": quarter_data,
        "annual_totals": annual_totals,
        "usd_totals": usd_totals,
        "year": year,
        "extension": ext,
        "currency_symbol": sym,
        "currency_code": ext_currency.code if ext_currency else "Non renseignée",
        "exchange_rate_to_usd": ext_currency.usd_rate if ext_currency else 1,
        "show_usd": bool(ext_currency and ext_currency.code != "USD" and not is_consolidated),
        "is_consolidated": is_consolidated,
        "extensions": extensions,
        "is_admin": user_is_admin(request.user),
        "extras_income": extras_income,
        "extras_expense": extras_expense,
    })


# ─────────────────────────────────────────────────────────────────────────────
# Helpers internes
# ─────────────────────────────────────────────────────────────────────────────

def _aggregate_totals(reports, year, month=None, months=None, user=None, extension_id=None):
    from apps.churches.models import ChurchExtension
    from .models import ExtraIncome, ExtraExpense, ServiceReport

    # Normaliser: accepter une liste de rapports (vues trimestrielles/annuelles)
    # en la convertissant en queryset pour supporter .filter()/.count()/.select_related()
    if isinstance(reports, (list, tuple)):
        reports = ServiceReport.objects.filter(pk__in=[r.pk for r in reports])

    # 1. Déterminer l'extension sélectionnée (si admin) ou imposée (si utilisateur d'extension)
    selected_ext = None
    if user and not user_is_admin(user):
        selected_ext = user_extension(user)
    elif extension_id:
        try:
            selected_ext = ChurchExtension.objects.get(pk=int(extension_id))
        except (ChurchExtension.DoesNotExist, ValueError, TypeError):
            pass

    # 2. Filtrer les rapports de culte
    reports_filtered = reports
    if selected_ext:
        reports_filtered = reports.filter(extension=selected_ext)

    # 3. Filtrer les recettes/dépenses supplémentaires
    extras_income_qs = ExtraIncome.objects.filter(income_date__year=year)
    extras_expense_qs = ExtraExpense.objects.filter(expense_date__year=year)
    
    if month:
        extras_income_qs = extras_income_qs.filter(income_date__month=month)
        extras_expense_qs = extras_expense_qs.filter(expense_date__month=month)
    if months:
        extras_income_qs = extras_income_qs.filter(income_date__month__in=months)
        extras_expense_qs = extras_expense_qs.filter(expense_date__month__in=months)
        
    if selected_ext:
        extras_income_qs = extras_income_qs.filter(extension=selected_ext)
        extras_expense_qs = extras_expense_qs.filter(extension=selected_ext)
    elif user and not user_is_admin(user):
        ext = user_extension(user)
        if ext:
            extras_income_qs = extras_income_qs.filter(extension=ext)
            extras_expense_qs = extras_expense_qs.filter(extension=ext)

    if user and not user_is_admin(user) and not selected_ext:
        reports_filtered = reports_filtered.none()
        extras_income_qs = extras_income_qs.none()
        extras_expense_qs = extras_expense_qs.none()
    is_consolidated = (user_is_admin(user) and not selected_ext)
    
    target = Currency.get_default() if is_consolidated or not selected_ext else get_currency_for_extension(selected_ext)
    fields = {
        "total_offrandes": "total_offerings", "total_dimes": "tithe_deduction",
        "total_social": "social_deduction", "total_depenses": "total_expenses",
        "total_reste_culte": "net_balance",
    }
    totals = {key: Decimal("0.00") for key in fields}
    totals.update(total_cultes=0, total_presence=0)
    # Evaluate THIS queryset so display annotations survive when templates iterate it.
    reports_filtered = reports_filtered.select_related("currency", "extension__currency")
    for report in reports_filtered:
        source = get_record_currency(report)
        totals["total_cultes"] += 1
        totals["total_presence"] += report.total_attendance
        for key, field in fields.items():
            amount = convert_currency(getattr(report, field), source, target)
            totals[key] += amount
            setattr(report, field + "_disp", amount)

    def extras(queryset, date_field):
        total = Decimal("0.00")
        rows = []
        for item in queryset.select_related("currency", "extension__currency"):
            amount = convert_currency(item.amount, get_record_currency(item), target)
            total += amount
            rows.append({"date": getattr(item, date_field), "description": item.description,
                         "amount": amount, "extension": item.extension.name if is_consolidated else "",
                         "currency": target.symbol})
        return total, rows

    totals["total_extras"], extras_income_list = extras(extras_income_qs, "income_date")
    totals["total_extra_expenses"], extras_expense_list = extras(extras_expense_qs, "expense_date")
    totals["total_reste"] = money(totals["total_reste_culte"] + totals["total_extras"] - totals["total_extra_expenses"])
    return totals, selected_ext, is_consolidated, reports_filtered, extras_income_list, extras_expense_list


def _totals_in_usd(totals, ext):
    """Convertit un dictionnaire de totaux en USD via le système Currency."""
    from apps.churches.models import Currency

    if not ext:
        return None
    from_currency = get_currency_for_extension(ext)
    usd = Currency.objects.filter(code="USD", is_active=True).first()
    if not from_currency or not usd or from_currency == usd:
        return None
    non_monetary = {"total_cultes", "total_presence"}
    return {
        k: (v if k in non_monetary else convert_currency(v, from_currency, usd))
        for k, v in totals.items()
    }
